import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

F="Calibri Light"; MONO="Consolas"
doc=Document()

def setfont(obj,name):
    rpr=obj.get_or_add_rPr() if hasattr(obj,'get_or_add_rPr') else obj
    rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
    for a in list(rf.attrib):
        if 'Theme' in a: del rf.attrib[a]
    for a in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'): rf.set(qn(a),name)

# ---- styles
st=doc.styles['Normal']; st.font.name=F; st.font.size=Pt(11)
setfont(st.element.get_or_add_rPr(),F)
pf=st.paragraph_format; pf.line_spacing=1.25; pf.space_after=Pt(6); pf.space_before=Pt(0)
pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name,size,align,before,after in [('Heading 1',15,WD_ALIGN_PARAGRAPH.CENTER,0,12),('Heading 2',13,WD_ALIGN_PARAGRAPH.LEFT,12,6),('Heading 3',12,WD_ALIGN_PARAGRAPH.LEFT,10,4)]:
    h=doc.styles[name]; h.font.name=F; h.font.size=Pt(size); h.font.bold=True; h.font.italic=False
    h.font.color.rgb=RGBColor(0,0,0); setfont(h.element.get_or_add_rPr(),F)
    h.paragraph_format.alignment=align; h.paragraph_format.space_before=Pt(before); h.paragraph_format.space_after=Pt(after)
    h.paragraph_format.line_spacing=1.2; h.paragraph_format.keep_with_next=True
doc.styles['Heading 1'].paragraph_format.page_break_before=False
for nm in ('List Bullet','Caption','Title'):
    try:
        x=doc.styles[nm]; x.font.name=F; setfont(x.element.get_or_add_rPr(),F); x.font.color.rgb=RGBColor(0,0,0)
    except KeyError: pass
cap=doc.styles['Caption']; cap.font.size=Pt(10.5); cap.font.italic=True; cap.font.bold=False
cap.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap.paragraph_format.space_after=Pt(10); cap.paragraph_format.line_spacing=1.15
from docx.enum.style import WD_STYLE_TYPE
code=doc.styles.add_style('Code Block',WD_STYLE_TYPE.PARAGRAPH); code.base_style=st
code.font.name=MONO; code.font.size=Pt(8.5); setfont(code.element.get_or_add_rPr(),MONO)
cpf=code.paragraph_format; cpf.line_spacing=1.0; cpf.space_after=Pt(0); cpf.alignment=WD_ALIGN_PARAGRAPH.LEFT; cpf.keep_with_next=True
cpf.left_indent=Cm(0.3)
ppr=code.element.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),'F2F2F2'); ppr.append(shd)

for lvl in (1,2,3):
    ts=doc.styles.add_style('toc %d'%lvl,WD_STYLE_TYPE.PARAGRAPH); ts.base_style=st
    ts.paragraph_format.line_spacing=1.15; ts.paragraph_format.space_after=Pt(3 if lvl>1 else 4); ts.paragraph_format.space_before=Pt(6 if lvl==1 else 0)
    ts.paragraph_format.left_indent=Cm(0.6*(lvl-1)); ts.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    ts.font.size=Pt(11.5 if lvl==1 else 11); ts.font.bold=(lvl==1)
# ---- page setup
sec=doc.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.left_margin=Cm(3.0); sec.right_margin=Cm(2.5); sec.top_margin=Cm(2.5); sec.bottom_margin=Cm(2.5)

PH=re.compile(r'(\[Confirm/insert:[^\]]+\])')
def addruns(p,text,size=None,bold=False,italic=False):
    for part in PH.split(text):
        if not part: continue
        r=p.add_run(part); r.bold=bold; r.italic=italic
        if size: r.font.size=Pt(size)
        if PH.fullmatch(part): r.font.highlight_color=WD_COLOR_INDEX.YELLOW
    return p
def para(text="",align=None,size=None,bold=False,italic=False,after=None,before=None,style=None,ls=None):
    p=doc.add_paragraph(style=style)
    if text: addruns(p,text,size,bold,italic)
    if align is not None: p.alignment=align
    if after is not None: p.paragraph_format.space_after=Pt(after)
    if before is not None: p.paragraph_format.space_before=Pt(before)
    if ls is not None: p.paragraph_format.line_spacing=ls
    return p
def field(p,instr,placeholder=""):
    r=p.add_run(); f=OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'),'begin'); f.set(qn('w:dirty'),'true'); r._r.append(f)
    r=p.add_run(); t=OxmlElement('w:instrText'); t.set(qn('xml:space'),'preserve'); t.text=instr; r._r.append(t)
    r=p.add_run(); f=OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'),'separate'); r._r.append(f)
    if placeholder: p.add_run(placeholder)
    r=p.add_run(); f=OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'),'end'); r._r.append(f)
def pagebreak():
    p=doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE); p.paragraph_format.space_after=Pt(0)
def frontheading(text):
    p=para(text,align=WD_ALIGN_PARAGRAPH.CENTER,size=15,bold=True,after=12)
    p.paragraph_format.page_break_before=True
    return p
def pgnum(section,fmt,start=None):
    sp=section._sectPr; pg=sp.find(qn('w:pgNumType'))
    if pg is None: pg=OxmlElement('w:pgNumType'); sp.append(pg)
    pg.set(qn('w:fmt'),fmt)
    if start: pg.set(qn('w:start'),str(start))
def footer_num(section):
    section.footer.is_linked_to_previous=False
    fp=section.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in list(fp.runs): r._r.getparent().remove(r._r)
    field(fp,'PAGE','1')


# ================= TITLE PAGE
para("KUMASI TECHNICAL UNIVERSITY",WD_ALIGN_PARAGRAPH.CENTER,18,True,after=2,before=10)
para("[Confirm/insert: Faculty name]",WD_ALIGN_PARAGRAPH.CENTER,12,after=0)
para("[Confirm/insert: Department name]",WD_ALIGN_PARAGRAPH.CENTER,12,after=30)
para("INDUSTRIAL ATTACHMENT PROJECT REPORT",WD_ALIGN_PARAGRAPH.CENTER,13,True,after=10)
para("ACCOLM CHECK",WD_ALIGN_PARAGRAPH.CENTER,26,True,after=4)
para("A web tool that screens how websites collect personal data against Ghana's Data Protection Act, 2012 (Act 843)",WD_ALIGN_PARAGRAPH.CENTER,13,italic=True,after=6)
para("A solution to a problem I encountered during my industrial attachment with Accolm",WD_ALIGN_PARAGRAPH.CENTER,12,after=34)
t=doc.add_table(rows=0,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
for k,v in [("Name","Ernest Abel Nanabanyin Mensah"),("Student Number","052341360055"),("Programme","BTech Computer Technology"),("Attachment","Accolm (remote), 11 May – 30 September 2026"),("Academic Supervisor","[Confirm/insert: name]"),("Work-Based Supervisor","[Confirm/insert: name]")]:
    c=t.add_row().cells; c[0].width=Cm(5.0); c[1].width=Cm(9.5)
    for cell,txt,b in ((c[0],k,True),(c[1],v,False)):
        pp=cell.paragraphs[0]; addruns(pp,txt,12,b); pp.alignment=WD_ALIGN_PARAGRAPH.LEFT; pp.paragraph_format.line_spacing=1.15; pp.paragraph_format.space_after=Pt(3)
para("",after=40)
para("OCTOBER 2026",WD_ALIGN_PARAGRAPH.CENTER,14,True,before=40)

# ================= ABSTRACT
frontheading("ABSTRACT")
para("During my remote industrial attachment with Accolm (11 May to 30 September 2026) I trained in HTML, CSS, Bootstrap, JavaScript and SQL with MySQL. I learnt that web forms are where personal data enters a system, that browser-side checks are not security, and that Ghana's Data Protection Act, 2012 (Act 843) sets eight principles for anyone who processes personal data. However, I had no easy way to check whether a real website met even the basics of the Act, and many small Ghanaian businesses have the same gap.")
para("This report describes Accolm Check, a web application I built to close that gap. A user pastes a web address; the server fetches the page safely, finds the forms and fields that collect personal data (including Ghana Card, mobile money and digital address fields), runs fourteen checks, and maps each finding to a principle in section 17 of the Act with a plain-language explanation and a fix. The result is a score out of 100 and a shareable, printable report. The app uses HTML, CSS, Bootstrap 5 and vanilla JavaScript on the front end, Node.js with Express on the server, and SQLite or MySQL with prepared statements for scan history. It is protected against Server-Side Request Forgery, and 27 automated tests passed. The tool is an automated first-pass screening, not legal advice.")

# ================= TOC
frontheading("TABLE OF CONTENTS")
p=doc.add_paragraph(); p.paragraph_format.line_spacing=1.15
field(p,'TOC \\o "1-2" \\h \\z \\u','Right-click and choose Update Field to build the table of contents.')

footer_num(sec); sec.different_first_page_header_footer=True; pgnum(sec,'lowerRoman',1)
bodysec=doc.add_section(WD_SECTION.NEW_PAGE); pgnum(bodysec,'decimal',1); footer_num(bodysec); bodysec.different_first_page_header_footer=False

# ================= BODY
lines=open('report_content.md').read().split('\n')
i=0; first_h1=True
def make_table(rows):
    hdr=[c.strip() for c in rows[0].split('|')]
    t=doc.add_table(rows=1,cols=len(hdr)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,h in enumerate(hdr):
        cell=t.rows[0].cells[j]; pp=cell.paragraphs[0]; addruns(pp,h,9.5,True); pp.alignment=WD_ALIGN_PARAGRAPH.LEFT; pp.paragraph_format.line_spacing=1.15; pp.paragraph_format.space_after=Pt(2); pp.paragraph_format.keep_with_next=True
        tcPr=cell._tc.get_or_add_tcPr(); s=OxmlElement('w:shd'); s.set(qn('w:val'),'clear'); s.set(qn('w:color'),'auto'); s.set(qn('w:fill'),'D9E2F3'); tcPr.append(s)
    trPr=t.rows[0]._tr.get_or_add_trPr(); th=OxmlElement('w:tblHeader'); trPr.append(th)
    for r in rows[1:]:
        cs=[c.strip() for c in r.split('|')]; cells=t.add_row().cells
        for j,v in enumerate(cs):
            pp=cells[j].paragraphs[0]; addruns(pp,v,9.5); pp.alignment=WD_ALIGN_PARAGRAPH.LEFT; pp.paragraph_format.line_spacing=1.15; pp.paragraph_format.space_after=Pt(2); pp.paragraph_format.keep_with_next=True
    W={3:[1.9,6.4,7.2],2:[4.6,10.9],4:None}
    hdrkey=rows[0]
    widths={'Section':[1.9,6.4,7.2],'Risk':[4.6,10.9],'Test file':[3.3,1.2,8.6,2.4],'Site':[7.2,1.6,2.8,3.9]}.get(hdr[0])
    if widths:
        t.autofit=False
        tbl=t._tbl; grid=tbl.tblGrid
        for j,gc in enumerate(grid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int(widths[j]*567)))
        tblPr=tbl.tblPr; lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tblPr.append(lay)
        tw=tblPr.find(qn('w:tblW'))
        if tw is None: tw=OxmlElement('w:tblW'); tblPr.append(tw)
        tw.set(qn('w:w'),str(int(sum(widths)*567))); tw.set(qn('w:type'),'dxa')
        for row in t.rows:
            for j,cell in enumerate(row.cells): cell.width=Cm(widths[j])
    return t
while i<len(lines):
    L=lines[i]
    if not L.strip(): i+=1; continue
    if L.startswith('H1: '):
        h=doc.add_heading(L[4:],level=1)
        if L[4:]=='REFERENCES': h.paragraph_format.page_break_before=True
        if first_h1: h.paragraph_format.page_break_before=False; first_h1=False
        i+=1; continue
    if L.startswith('H2: '): doc.add_heading(L[4:],level=2); i+=1; continue
    if L.startswith('H3: '): doc.add_heading(L[4:],level=3); i+=1; continue
    if L.startswith('TABLE: '):
        rows=[L[7:]]; i+=1
        while i<len(lines) and lines[i].strip() and not lines[i].startswith('CAPTION:'): rows.append(lines[i]); i+=1
        make_table(rows); continue
    if L.startswith('IMG: '):
        path,w=[x.strip() for x in L[5:].split('|')]
        pp=doc.add_paragraph(); pp.alignment=WD_ALIGN_PARAGRAPH.CENTER; pp.paragraph_format.keep_with_next=True
        pp.paragraph_format.line_spacing=1.0; pp.paragraph_format.space_before=Pt(6); pp.paragraph_format.space_after=Pt(2)
        pp.add_run().add_picture(path,width=Cm(float(w))); i+=1; continue
    if L.startswith('IMG2: '):
        a,wa,b,wb=[x.strip() for x in L[6:].split('|')]
        t=doc.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        for cell,path,w in ((t.rows[0].cells[0],a,wa),(t.rows[0].cells[1],b,wb)):
            cell.width=Cm(float(w)+0.4); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            pp=cell.paragraphs[0]; pp.alignment=WD_ALIGN_PARAGRAPH.CENTER; pp.paragraph_format.line_spacing=1.0; pp.paragraph_format.space_after=Pt(0); pp.paragraph_format.keep_with_next=True
            pp.add_run().add_picture(path,width=Cm(float(w)))
        trPr=t.rows[0]._tr.get_or_add_trPr(); cs=OxmlElement('w:cantSplit'); trPr.append(cs)
        i+=1; continue
    if L.startswith('CAPTION: '):
        para(L[9:],style='Caption',before=4); i+=1; continue
    if L.startswith('CODE:'):
        i+=1; block=[]
        while i<len(lines) and not lines[i].startswith('CAPTION:'): block.append(lines[i]); i+=1
        while block and not block[-1].strip(): block.pop()
        para("",after=0,ls=1.0)
        for k,b in enumerate(block):
            pp=doc.add_paragraph(style='Code Block'); pp.add_run(b if b else " ")
            if k==len(block)-1: pp.paragraph_format.keep_with_next=True
        continue
    if L.startswith('REF: '):
        pp=para(L[5:],align=WD_ALIGN_PARAGRAPH.LEFT,after=6,ls=1.15)
        pp.paragraph_format.left_indent=Cm(1.27); pp.paragraph_format.first_line_indent=Cm(-1.27)
        i+=1; continue
    if L.startswith('- '):
        pp=para(L[2:],style='List Bullet',after=3); pp.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; i+=1; continue
    para(L); i+=1

# Settings: ask Word to refresh fields on open is not set (LibreOffice pre-fills TOC).
doc.save('draft.docx')
print("ok")
